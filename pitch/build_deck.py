"""Builds pitch/OpEx_AI_Builder_Cup_2026.pptx from the official AI Builder Cup template (pitch/template.pptx,
downloaded from the Hack2skill Google Slides link on 18 Oct 2026, plus two copies of its blank slide).

The template's header banner, footer strip and section titles are left as they are; our content goes underneath.
Numbers come from pitch/numbers.json (the app's default 5-year simulation run, seed 2026) and the model score
files in models/, so the deck cannot disagree with the product.

    python pitch/build_deck.py      (needs python-pptx and Pillow)
"""
import copy
import json
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PITCH = ROOT / "pitch"
SHOTS = PITCH / "screenshots"
CROPS = PITCH / "assets" / "crops"
N = json.loads((PITCH / "numbers.json").read_text())
FM = json.loads((ROOT / "models" / "failure_mode_metrics.json").read_text())
RUL = json.loads((ROOT / "models" / "rul_metrics.json").read_text())
PRE = json.loads((ROOT / "models" / "precursor_metrics.json").read_text())
D, K, C = N["delta"], N["kpis"], N["costs"]
SENS = json.loads((PITCH / "sensitivity.json").read_text())  # [adoption, coverage, net Cr, roi, payback, breakdowns]

INK, MUTED, RULE, LIGHT = "202124", "5F6368", "DADCE0", "F1F3F4"
BLUE, BLUE_TINT, RED, GREEN = "1A56DB", "E8F0FE", "B3261E", "1E7D4A"
FONT = "Arial"
FX = 95  # rupees per US dollar, interbank close 17 Oct 2026 (about 94.7 to 95.0)

TOP = 1.32          # first free line under the template's section title
LEFT, WIDTH = 0.39, 9.25
BOTTOM = 5.42       # template footer strip starts at 5.55


def rgb(h):
    return RGBColor.from_string(h)


def lakh(v):
    return f"₹{v / 1e5:,.1f} L" if abs(v) < 1e7 else f"₹{v / 1e7:,.2f} Cr"


# ----------------------------------------------------------------------------- helpers
def text(slide, x, y, w, h, paragraphs, size=10, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=3, line=None, name=None):
    """paragraphs: str | list of (str | list of (text, style) runs). style keys: bold, color, size, italic."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    if name:
        box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(spacing)
        if line:
            p.line_spacing = line
        runs = para if isinstance(para, list) else [(para, {})]
        for content, style in runs:
            r = p.add_run()
            r.text = content
            r.font.name = FONT
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", bold)
            r.font.italic = style.get("italic", False)
            r.font.color.rgb = rgb(style.get("color", color))
    return box


def bullets(slide, x, y, w, h, items, size=10, color=INK, gap=3, indent=0.15, line=None):
    """Real PowerPoint bullets (buChar), not a typed dot."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        if line:
            p.line_spacing = line
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(indent)))
        pPr.set("indent", str(-Inches(indent)))
        bu = pPr.makeelement(qn("a:buChar"), {"char": "•"})
        pPr.append(bu)
        runs = item if isinstance(item, list) else [(item, {})]
        for content, style in runs:
            r = p.add_run()
            r.text = content
            r.font.name = FONT
            r.font.size = Pt(style.get("size", size))
            r.font.bold = style.get("bold", False)
            r.font.color.rgb = rgb(style.get("color", color))
    return box


def rect(slide, x, y, w, h, fill="FFFFFF", line=RULE, width=0.75, rounded=False, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        s.adjustments[0] = 0.06
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(width)
    else:
        s.line.fill.background()
    s.shadow.inherit = False
    s.text_frame.text = ""
    if name:
        s.name = name
    return s


def labelled_box(slide, x, y, w, h, title, sub=None, fill="FFFFFF", line=RULE, title_color=INK, size=9.5, sub_size=None):
    rect(slide, x, y, w, h, fill=fill, line=line, rounded=True)
    paras = [[(title, {"bold": True, "color": title_color, "size": size})]]
    if sub:
        paras.append([(sub, {"color": MUTED, "size": sub_size or size - 1.5})])
    text(slide, x + 0.08, y + 0.05, w - 0.16, h - 0.1, paras, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, spacing=1)


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=1.25):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    tail = ln.line._get_or_add_ln()
    tail.append(tail.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"}))
    return ln


def hline(slide, x, y, w, color=RULE, width=0.75):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + w), Inches(y))
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    return ln


def crop(file, left=500, top=150, right=None, bottom=None):
    """Cut a region out of a 2x screenshot (the app's sidebar ends at x=480)."""
    CROPS.mkdir(parents=True, exist_ok=True)
    out = CROPS / f"{Path(file).stem}_{left}_{top}_{right or 'r'}_{bottom or 'b'}.png"
    if not out.exists():
        with Image.open(SHOTS / file) as im:
            im.crop((left, top, right or im.width, bottom or im.height)).save(out)
    return out


def picture(slide, path, x, y, w, h, border=True):
    """Fit the image inside the box, keeping its aspect ratio, top-left aligned."""
    with Image.open(path) as im:
        ratio = im.width / im.height
    pw, ph = (h * ratio, h) if w / h > ratio else (w, w / ratio)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(pw), Inches(ph))
    if border:
        pic.line.color.rgb = rgb(RULE)
        pic.line.width = Pt(0.75)
    return pic


def _cell_border(cell, side, color=None, w=0.5):
    tcPr = cell._tc.get_or_add_tcPr()
    tag = qn(f"a:ln{side}")
    for e in tcPr.findall(tag):
        tcPr.remove(e)
    ln = tcPr.makeelement(tag, {"w": str(int(w * 12700)), "cap": "flat", "cmpd": "sng", "algn": "ctr"})
    if color:
        fill = ln.makeelement(qn("a:solidFill"), {})
        fill.append(fill.makeelement(qn("a:srgbClr"), {"val": color}))
    else:
        fill = ln.makeelement(qn("a:noFill"), {})
    ln.append(fill)
    # ln* elements must precede the cell fill in tcPr
    order = ["L", "R", "T", "B"]
    pos = sum(1 for s in order[:order.index(side)] if tcPr.find(qn(f"a:ln{s}")) is not None)
    tcPr.insert(pos, ln)


def table(slide, x, y, w, rows, col_widths, size=8.5, row_h=0.3, header_fill=LIGHT, bold_first_col=False,
          align_right_cols=(), pad=0.05):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(row_h * len(rows)))
    tbl = shape.table
    tbl.first_row = True
    tbl.horz_banding = False
    for i, cw in enumerate(col_widths):
        tbl.columns[i].width = Inches(cw)
    for r, row in enumerate(rows):
        tbl.rows[r].height = Inches(row_h)
        for c, value in enumerate(row):
            cell = tbl.cell(r, c)
            cell.margin_left = cell.margin_right = Inches(pad)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(header_fill if r == 0 else "FFFFFF")
            for side in "LRT":
                _cell_border(cell, side, None)
            _cell_border(cell, "B", RULE if r < len(rows) - 1 else "9AA0A6", 0.5 if r else 1.0)
            tf = cell.text_frame
            tf.word_wrap = True
            runs = value if isinstance(value, list) else [(str(value), {})]
            p = tf.paragraphs[0]
            if c in align_right_cols:
                p.alignment = PP_ALIGN.RIGHT
            for content, style in runs:
                run = p.add_run()
                run.text = content
                run.font.name = FONT
                run.font.size = Pt(style.get("size", size))
                run.font.bold = style.get("bold", r == 0 or (bold_first_col and c == 0))
                run.font.color.rgb = rgb(style.get("color", INK if r else MUTED))
    return shape


def style_chart(chart, colors, legend=True, label_fmt='0.0"%"', size=8):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
    plot = chart.plots[0]
    plot.gap_width = 60
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(size - 0.5)
    plot.data_labels.font.color.rgb = rgb(INK)
    plot.data_labels.number_format = label_fmt
    plot.data_labels.number_format_is_linked = False
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    for series, color in zip(plot.series, colors):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = rgb(color)
        series.format.line.fill.background()
    chart.value_axis.visible = False
    chart.value_axis.has_major_gridlines = False
    chart.category_axis.tick_labels.font.size = Pt(size)
    chart.category_axis.tick_labels.font.color.rgb = rgb(INK)
    chart.category_axis.format.line.color.rgb = rgb(RULE)


def footnote(slide, words, y=5.2):
    text(slide, LEFT, y, WIDTH, 0.3, words, size=7.5, color=MUTED, spacing=0)


def headline(slide, words, y=TOP, size=13, h=0.5):
    text(slide, LEFT, y, WIDTH, h, words, size=size, bold=True, color=INK, name="Headline")


def notes(slide, words):
    slide.notes_slide.notes_text_frame.text = words


def title_box(slide):
    return next(s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip() and s.top < Inches(1))


def set_title(slide, title):
    """Keep the template's title box, font and size; replace its text."""
    tf = title_box(slide).text_frame
    first = tf.paragraphs[0]
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    runs = first.runs
    runs[0].text = title
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def give_title(blank, donor, title):
    """Blank template slide: copy the donor slide's title box so the font matches, then set the text."""
    src = title_box(donor)
    el = copy.deepcopy(src._element)
    blank.shapes._spTree.append(el)
    set_title(blank, title)


# ----------------------------------------------------------------------------- slides
def slide_team(s):
    box = next(sh for sh in s.shapes if sh.has_text_frame and "Team Details" in sh.text_frame.text)
    box.width = Inches(6.3)
    values = {"Team name: ": "OpEx", "Team leader name: ": "Srihitha Pulapa",
              "Problem Statement: ": "Manufacturing: Intelligent Operations & Industrial Efficiency"}
    for p in box.text_frame.paragraphs:
        label = "".join(r.text for r in p.runs)
        if label in values:
            r = copy.deepcopy(p.runs[0]._r)
            p._p.append(r)
            p.runs[-1].text = values[label]
            p.runs[-1].font.bold = True
    logo = s.shapes.add_picture(str(PITCH / "assets" / "opex-logo.png"), Inches(6.85), Inches(3.3), Inches(2.6))
    logo.name = "OpEx logo"
    text(s, 6.55, 4.42, 3.2, 0.9, ["Operations copilot for CNC and auto-component plants.",
                                   "Finds the fault early, prices the fix in rupees, and an engineer decides.",
                                   "Built from scratch in the Cup's build window, 7 Sept to 18 Oct 2026."],
         size=8.5, color=MUTED, align=PP_ALIGN.CENTER, spacing=2)
    notes(s, "Team OpEx, Manufacturing track. Working prototype: github.com/srihitha2005/Opex.")


def slide_brief(s):
    headline(s, "A plant usually finds out about a failing spindle bearing when the machine stops. OpEx finds out "
                "hours to days earlier, prices every option in rupees, and an engineer makes the call.", h=0.6)
    paras = [
        [("Who it is for. ", {"bold": True}),
         ("Tier-2 and Tier-3 auto-component plants with 10 to 50 CNC machines. ACMA has 1,100+ member companies of this "
          "kind; the industry turned over ₹7.59 lakh crore in FY26 and employs 1.5 million people directly.¹", {})],
        [("The problem. ", {"bold": True}),
         ("88% of Indian industrial businesses have an unplanned outage at least once a month and 19% still run "
          "machines until they fail.² Alarm screens list fifty thresholds, not one cause, and nobody knows what a fix "
          "costs until the invoice arrives. In our plant model a machine-hour is worth ₹4,200 to 6,500 of lost "
          "contribution, so one bearing failure with secondary damage costs ₹6 to 9 lakh.", {})],
        [("What OpEx does. ", {"bold": True}),
         ("It reads every machine each plant minute, turns abnormal signals into one plain-language cause with a failure "
          "window, prices seven ways of dealing with it, has three Gemini agents review the case, and records whether the "
          "chosen action worked. Operators can ask it anything, by voice, in English or five Indian languages.", {})],
    ]
    text(s, LEFT, 2.02, 5.1, 2.4, paras, size=9.5, spacing=6, line=1.06)
    hline(s, LEFT, 4.5, 5.1)
    text(s, LEFT, 4.6, 5.1, 0.45, [[
        ("Simulated, 24 machines, 5 years: ", {"bold": True}),
        (f"{D['breakdowns_reduction']:.0%} fewer breakdowns, {lakh(D['net_benefit_inr'])} net after paying for OpEx, "
         f"{D['roi_multiple']:.1f}× the cost. Arithmetic and sensitivity on the slide 'How we calculated the 5-year numbers'.", {})]],
        size=8.5, spacing=0, line=1.05)
    picture(s, crop("02_overview.png", 500, 370, 2880, 1800), 5.7, 2.02, 4.0, 2.55)
    text(s, 5.7, 4.62, 4.0, 0.5, "The overview as the engineer sees it: M-17 first, failure expected in 5 to 8 h, then three "
         "'plan it' items. Ranked by chance of failure × cost × urgency.", size=8, color=MUTED, spacing=0)
    footnote(s, "¹ ACMA, Industry Performance Review FY2025-26. ² ABB, Value of Reliability survey, Sapio Research, July 2023, n = 3,215, India cut. "
                "Plant economics are our model's assumptions, editable in the app; every number is listed on the Sources slide.", y=5.2)
    notes(s, "Open with the M-17 story on the right; the rest of the deck shows how that one line becomes a priced, approved decision.")


def slide_opportunities(s):
    set_title(s, "Opportunities")
    qa = [
        ("How different is it from any of the other existing ideas?",
         "Most predictive-maintenance tools stop at 'risk: high' and a work order. OpEx carries on to the decision: one root "
         "cause across machine, quality and energy signals; every option priced in ₹ with its chance of still breaking down; and "
         "a record of how often each action has worked in this plant, which re-ranks the options next time. When sensor data is "
         "bad it says so and lowers the stated confidence, and a drifting sensor is diagnosed as a sensor fault."),
        ("How will it be able to solve the problem?",
         "Warnings arrive about a day ahead on average (median lead 24 to 32 h on held-out Azure PdM data; the M-17 example is a "
         "fast failure caught 5 to 8 h out), so the part is on the shelf and the swap lands in the Sunday window instead of at "
         "2 a.m. The engineer sees downtime, cost, failure risk and extra scrap side by side and approves in one screen. The "
         "plant head gets a five-year with-versus-without case before buying a single sensor."),
        ("USP of the proposed solution",
         "A Gemini incident room in which a Diagnostician, a Planner and a Safety officer argue the case from measured facts, "
         "technicians' past notes and the plant's own SOPs, in the user's language. 'Rewind the factory', which replays one of "
         "the plant's own breakdowns with OpEx watching. And a simulator that charges itself for missed detections, false alarms "
         "and ignored advice, then shows how the answer moves when those assumptions get worse."),
    ]
    y = TOP
    for q, a in qa:
        text(s, LEFT, y, 5.0, 1.3, [[(q, {"bold": True, "size": 10})], [(a, {"size": 9})]], spacing=3, line=1.06)
        y += 1.34
    picture(s, crop("05_decide.png", 540, 320, 2880, 1270), 5.55, TOP, 4.15, 2.0)
    text(s, 5.55, 3.12, 4.15, 0.45, "Options for M-17: downtime, ₹ expected cost, failure risk, extra bad parts, energy, 'worked here before'. "
         "Unsafe options stay visible but are never recommended.", size=7.5, color=MUTED, spacing=0)
    picture(s, crop("07_sensor_drift.png", 540, 860, 1640, 1200), 5.55, 3.62, 4.15, 1.3)
    text(s, 5.55, 4.95, 4.15, 0.3, "M-21: vibration rising but temperature and power normal, 31% of readings missing, so the "
         "diagnosis is sensor drift and the confidence drops from 76% to 60%.", size=7.5, color=MUTED, spacing=0)
    notes(s, "Answer the three template questions in order; point at the 'Worked here before' column and the 'Unsafe' tags on the right.")


def slide_features(s):
    rows = [["Feature", "What you see in the app", "Who"],
            ["Ranked factory health", "OEE 75% = 83% availability × 92% speed × 97.8% quality. 'What should I worry about first?' orders "
             "problems by chance of failure × cost if it fails × how soon, so one machine is at the top, not fifty alarms", "Everyone"],
            ["Diagnosis in plain words", "'Vibration is 160% above its normal range (1.20 to 2.08 mm/s) and has been outside it for 180 minutes.' One root "
             "cause with a chain of impact and 'similar to 7 past incidents, 7 confirmed' from 194 incidents and 793 jobs", "Engineer"],
            ["Honest about data", "Missing readings, drift, calibration age and a physics check (power vs V × I) lower the stated confidence: "
             "'reduced from 76% to 60% because of these data problems'. A drifting sensor is called a sensor fault", "Engineer"],
            ["Seven priced options", "Stop, slow 10%, continue, shift load, partial fix, replace, watch. Each shows downtime, ₹ expected cost, failure "
             "risk, extra bad parts, energy and 'worked 6 of 6 times'. Unsafe options are flagged and never recommended", "Engineer"],
            ["Gemini incident room", "Diagnostician, Planner and Safety officer review the case from measured facts, past technician notes and the "
             "plant's SOPs (lockout-tagout steps included). Their reasoning is saved with the decision", "Engineer"],
            ["Ask OpEx", "Gemini with six function tools, so every number comes from plant data. Voice in, read-aloud, English plus five Indian languages, "
             "photo check of a finished part, a shift-handover note written from live data", "Operator"],
            ["Approval and learning", "Approve, reject or pick another option; roles enforced on the server; outcome re-checked 2 plant hours later and "
             "fed back into the ranking. A 5-year with-vs-without simulation and 'Rewind the factory' for the plant head", "All roles"]]
    table(s, LEFT, TOP + 0.02, WIDTH, rows, [1.65, 6.75, 0.85], size=8.8, row_h=0.43, bold_first_col=True)
    footnote(s, "Three different numbers appear next to a machine: chance of failure (the prediction model), pattern match (how closely the signals resemble a "
                "known fault) and confidence (lowered by data problems). The incident history is seeded from Azure PdM maintenance records re-dated onto this plant; "
                "a real plant starts from its own CMMS export. Screens in this deck were captured on different days of the live twin, so M-17's numbers vary.", y=4.93)
    notes(s, "Every row is working in the prototype; the quoted sentences are copied from the screens.")


def slide_flow(s):
    headline(s, "The top row runs on its own every plant minute. The bottom row always includes a person, and what the person "
                "decides changes the ranking next time.", h=0.6)
    text(s, LEFT, 1.98, 4, 0.2, "Automatic", size=8, color=MUTED)
    top_y, bw, gap = 2.16, 1.67, 0.2
    steps = [("Readings", "vibration, temperature, rpm,\ntorque, current, coolant pressure"),
             ("Data health", "missing, drift, physics check\n→ confidence multiplier"),
             ("Detect", "how far outside normal,\nfor how long"),
             ("Diagnose", "one cause, chain of impact,\n'seen this before?'"),
             ("Predict", "failure window, remaining\nlife, quality, energy")]
    x = LEFT
    for i, (t, sub) in enumerate(steps):
        labelled_box(s, x, top_y, bw, 0.78, f"{i + 1}  {t}", sub, size=9, sub_size=7.5)
        if i < 4:
            arrow(s, x + bw + 0.02, top_y + 0.39, x + bw + gap - 0.02, top_y + 0.39, BLUE)
        x += bw + gap
    arrow(s, LEFT + 4 * (bw + gap) + bw / 2, top_y + 0.8, LEFT + 4 * (bw + gap) + bw / 2, 3.42, BLUE)
    text(s, LEFT, 3.12, 4, 0.2, "With a person in the loop (read right to left)", size=8, color=MUTED)
    low_y = 3.44
    lower = [("10  Learn", "success counts per\nproblem × action", "E6F4EA", "9CCFB0"),
             ("9  Act, then check", "outcome measured 2 plant\nhours later", "FFFFFF", RULE),
             ("8  Human decides", "approve / reject / modify,\nby role, with a note", BLUE_TINT, BLUE),
             ("7  Gemini incident room", "Diagnostician · Planner ·\nSafety officer", "FFFFFF", RULE),
             ("6  Price 7 options", "₹ cost, failure risk, scrap,\nenergy, 'worked before'", "FFFFFF", RULE)]
    for k, (t, sub, fill, line) in enumerate(lower):
        px = LEFT + k * (bw + gap)
        labelled_box(s, px, low_y, bw, 0.78, t, sub, fill=fill, line=line, size=9, sub_size=7.5)
        if k > 0:
            arrow(s, px - 0.02, low_y + 0.39, px - gap + 0.02, low_y + 0.39, BLUE)
    ly = low_y + 0.78 + 0.16
    for (x1, y1, x2, y2) in [(LEFT + bw / 2, low_y + 0.8, LEFT + bw / 2, ly), (LEFT + bw / 2, ly, LEFT + 4 * (bw + gap) + bw / 2, ly)]:
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        ln.line.color.rgb = rgb(GREEN); ln.line.width = Pt(1.5)
    arrow(s, LEFT + 4 * (bw + gap) + bw / 2, ly, LEFT + 4 * (bw + gap) + bw / 2, low_y + 0.82, GREEN, 1.5)
    text(s, LEFT + bw / 2 + 0.15, ly + 0.03, 6, 0.22, "What worked in this plant feeds step 6 the next time the same problem appears",
         size=8.5, color=GREEN, bold=True)
    who = [("Operator", "acknowledges alerts, approves tool changes, standby and recalibration, asks by voice"),
           ("Maintenance engineer", "approves stops, speed cuts and part swaps; writes investigations; adds SOPs"),
           ("Plant manager", "sets priorities and ₹ assumptions, reads the learning log, runs the 5-year case")]
    for i, (w, what) in enumerate(who):
        text(s, LEFT + i * 3.1, 4.78, 2.95, 0.5, [[(w + ". ", {"bold": True, "size": 8.5}), (what, {"size": 8.5, "color": MUTED})]], spacing=0)
    footnote(s, "Warnings are held for 45 plant minutes so they do not flicker (ISA-18.2). Screens stay grey when normal; colour appears "
                "with an icon and a word only when something is wrong (ISA-101).", y=5.22)
    notes(s, "Steps 1 to 5 are automatic; 6 to 10 include a human decision; the green line is the learning loop.")


def slide_wireframes(s):
    headline(s, "The same four questions on a phone at the machine and on a desktop in the maintenance office")
    picture(s, crop("23_mobile_overview.png", 0, 0, 780, 1420), LEFT, 1.9, 1.76, 3.2)
    picture(s, crop("24_mobile_m17.png", 0, 0, 780, 1330), 2.3, 1.9, 1.88, 3.2)
    picture(s, crop("04_machine_top.png", 500, 150, 2880, 1800), 4.4, 1.9, 5.25, 2.55)
    text(s, 4.4, 4.55, 5.25, 0.7, "Every machine page answers, in order: 1 What we detected, 2 Why it is happening, 3 What happens next, "
         "4 What should we do. Jargon has a '?' with a plain explanation, long lists show the top three first, and nothing times out.",
         size=8.5, spacing=0)
    text(s, LEFT, 5.15, 3.4, 0.3, "Phone: the overview and M-17 during the same critical episode.", size=7.5, color=MUTED)
    notes(s, "Screenshots at phone and desktop widths. Status is always icon plus word plus colour, never colour alone.")


def slide_architecture(s):
    headline(s, "One Cloud Run container with Gemini beside it and Postgres underneath")
    labelled_box(s, LEFT, 1.95, 1.55, 0.85, "Shop floor", "CNC controllers and sensors\nvia OPC UA / MTConnect\ngateway or direct push", fill=LIGHT, size=9, sub_size=7.5)
    labelled_box(s, LEFT, 3.2, 1.55, 0.7, "Pub/Sub", "readings in,\ncritical alerts out", size=9, sub_size=7.5)
    arrow(s, LEFT + 0.78, 2.82, LEFT + 0.78, 3.18)
    fx, fy, fw, fh = 2.2, 1.8, 4.5, 2.78
    rect(s, fx, fy, fw, fh, fill="FFFFFF", line=BLUE, width=1.0)
    text(s, fx + 0.12, fy + 0.07, fw - 0.2, 0.25, [[("Cloud Run, one container, asia-south1 ", {"bold": True, "color": BLUE}),
                                                    ("· one URL for every device", {"color": MUTED, "size": 7.5})]], size=9.5, spacing=0)
    inner = [("React 19 web app", "accessible UI, voice,\n6 languages"), ("FastAPI", "Google sign-in, roles,\nAPI, ingest endpoint"),
             ("Live plant twin", "1-min readings, analytics\ncached per plant minute"), ("Decision engine", "7 priced options, ranking,\nlearning loop"),
             ("5-year simulator", "2.1 M machine-hours\nin about 4.5 s"), ("Gemini agents + copilot", "function calling,\nhybrid search over SOPs")]
    for i, (t, sub) in enumerate(inner):
        col, row = i % 3, i // 3
        labelled_box(s, fx + 0.12 + col * 1.44, fy + 0.4 + row * 1.14, 1.34, 1.02, t, sub, fill=LIGHT, line=None, size=8.5, sub_size=7.5)
    arrow(s, LEFT + 1.57, 3.45, fx - 0.02, 3.45)
    arrow(s, LEFT + 1.57, 2.37, fx - 0.02, 2.37)
    services = [("Gemini", "3.5-flash-lite default · 3.8-flash · embedding-001\nVertex AI in production, AI Studio key for the demo", BLUE_TINT, BLUE),
                ("Postgres: Supabase (Mumbai) today", "users, roles, decisions, incidents, runs, audit log · RLS on\nCloud SQL for PostgreSQL is a connection-string change", "FFFFFF", RULE),
                ("BigQuery sandbox", "daily simulation results, decision snapshot\nfor Looker Studio", "FFFFFF", RULE),
                ("Cloud Storage", "sensor CSV uploads, part photos", "FFFFFF", RULE)]
    for i, (t, sub, fill, line) in enumerate(services):
        y = 1.8 + i * 0.72
        labelled_box(s, 7.0, y, 2.65, 0.64, t, sub, fill=fill, line=line, size=8.5, sub_size=6.8)
        arrow(s, fx + fw + 0.02, y + 0.32 + 0.01 * i, 6.98, y + 0.32 + 0.01 * i)
    text(s, LEFT, 4.68, WIDTH, 0.5, [
        [("Scale path. ", {"bold": True}), ("The demo keeps one plant twin in memory on a single instance. With real readings, state lives in Postgres and "
          "Pub/Sub, so Cloud Run can run as many instances as a plant group needs; the pipeline that trains the three models and calibrates the simulator "
          "runs in 30 s and needs no extra service at this scale.", {})],
        [("If Gemini is unreachable ", {"bold": True}), ("the incident room, copilot, shift briefing, photo check and investigation writer fall back to "
          "rule-based text and are labelled as such on screen. Detection, prediction and pricing are deterministic and unaffected.", {})]],
        size=8, spacing=3, line=1.05)
    notes(s, "Single deployable unit. The right column is what Google Cloud and Supabase provide. The scale path answers 'max 1 instance'.")


def slide_tech(s):
    headline(s, "Gemini does the reasoning and the language. Deterministic code does the arithmetic, so every number can be audited.")
    rows = [["Technology", "What it does in OpEx", "Why this choice"],
            ["Gemini via google-genai: gemini-3.5-flash-lite default, gemini-3.8-flash optional. Vertex AI for production (one flag); AI Studio key for the free demo",
             "Three incident-room agents (structured JSON); copilot with 6 function tools; investigation writer; shift briefing; part-photo check; 6 languages",
             "Agents receive facts the engine already computed, so a light model is enough; its free tier covers a demo day"],
            ["gemini-embedding-001 + TF-IDF", "Hybrid search over factory memory: SOPs, manuals, 194 past incidents with technician notes", "About 250 chunks today; numpy in milliseconds, pgvector when it passes ~100k"],
            ["Cloud Run (asia-south1)", "One container: web app, API, live twin, analytics, simulator, agents", "One URL to judge, scales to zero, one gcloud command to deploy"],
            ["Supabase Postgres with Row Level Security", "Users, roles, decisions, outcomes, incidents, runs, audit log", "Managed Postgres in Mumbai; SQLite fallback for an offline laptop"],
            ["Google sign-in", "OAuth ID token verified on the server, httpOnly session cookie, roles checked on every request", "People use the account they have; no passwords stored for real users"],
            ["BigQuery sandbox · Pub/Sub · Cloud Storage", "Simulation and decision tables for Looker Studio; sensor push in and alerts out; uploads and photos", "Each has a concrete job and sits inside its free tier at demo scale"],
            ["Python 3.12, FastAPI, scikit-learn, numpy", "Three trained models, analytics, 5-year simulator with shared random draws", "11,400 lines of code; 38 automated tests pass"],
            ["React 19, TypeScript, Recharts, Atkinson Hyperlegible", "WCAG 2.2 AA interface, voice in and out, phone to desktop", "Shop floors have gloves, glare and older eyes"]]
    table(s, LEFT, 1.9, WIDTH, rows, [2.75, 3.6, 2.9], size=7.8, row_h=0.33, bold_first_col=True)
    footnote(s, "Not used at this scale: Dataflow and Vertex AI training (the models train in 30 s inside the container). The next slide lists what we would add "
                "with a budget and what each costs today.", y=5.24)
    notes(s, "Every service maps to a module in backend/app/integrations or backend/app/ai.")


def slide_tech_optional(s, donor):
    give_title(s, donor, "Technologies to be used in the solution (2 of 2): held back for cost")
    headline(s, "The prototype runs entirely on free tiers so any plant can try it before paying. These are the paid services we would "
                "add, what each improves, and its list price today.", y=1.2, size=11, h=0.5)
    rows = [["Service", "What it would add to OpEx", "Why it beats what runs today", "List price (USD)", "When"],
            ["Vertex AI for Gemini (already wired, off for the demo)", "Data residency in asia-south1, IAM, higher quotas, prompts never used for training", "The demo runs on a free AI Studio key; a plant's sensor data needs the enterprise terms", "Same per-token price as the API ($0.30 in / $2.50 out per 1M for 3.5 Flash-Lite); no free tier", "First paying plant"],
            ["gemini-3.8-flash as default", "Sharper multi-step reasoning in the incident room and investigations", "Flash-Lite was chosen for its free daily quota, not its depth", "$0.75 in / $3.75 out per 1M tokens, about 2.5× Flash-Lite", "About ₹8k a month"],
            ["Gemini Live API (3.8-live)", "Real two-way voice at the machine, interruptions, six languages", "Voice today uses the browser's speech engine, which varies by phone", "Audio $3.00 in / $12.00 out per 1M tokens (about $0.005 / $0.018 a minute)", "Operator pilot"],
            ["Speech-to-Text (Chirp 3)", "Server-side transcripts of operator notes in Indian languages", "Same quality on every device; transcript stored with the decision", "$0.016 a minute after 60 free minutes a month", "Operator pilot"],
            ["Vertex AI Agent Engine + ADK", "Hosted, traced, versioned agents with sessions and memory", "The three agents run in-process; tracing and evaluation are hand-rolled", "$0.0864 per vCPU-hour, $0.009 per GiB-hour after 50 free vCPU-hours", "Second plant"],
            ["Vertex AI Vector Search", "Factory memory beyond about 100k chunks", "In-memory search is enough for 250 chunks; pgvector is the step between", "About $0.75 per node-hour reported (about $550 a month minimum)", "Multi-plant memory"],
            ["Dataflow (streaming)", "Pub/Sub → BigQuery at more than 10k readings a second", "Readings go straight into the twin; BigQuery gets daily summaries", "About $0.069 per vCPU-hour plus $0.0036 per GiB-hour", "Past ~200 machines"],
            ["Visual Inspection AI", "Line-camera defect detection instead of a manual photo check", "Today an operator photographs a part and Gemini judges it", "$100 per camera stream per solution a month (restricted GA)", "Quality pilot"],
            ["Manufacturing Data Engine + Connect", "Native connectors for FANUC, Siemens and Mitsubishi controllers", "The prototype ingests through an API key; real plants need controller drivers", "Enterprise pricing via sales; a Kepware FANUC driver suite is about $995", "First real plant"]]
    table(s, LEFT, 1.78, WIDTH, rows, [1.6, 2.1, 2.1, 2.25, 1.2], size=7.2, row_h=0.33, bold_first_col=True)
    footnote(s, f"Public list prices read 17-18 Oct 2026 (ai.google.dev pricing, cloud.google.com pricing pages, Google Cloud release notes), ₹{FX} per USD; numbered on the Sources slide. "
                "Each row replaces a module that already has an interface in the code.", y=5.2)
    notes(s, "The honest answer to 'why not Vertex, why not Dataflow': money and scale. Each row names the trigger that would make us switch it on.")


def slide_cost(s):
    headline(s, f"₹{C['one_time_inr'] / 1e5:.1f} lakh to start and ₹{C['yearly_inr'] / 1e5:.1f} lakh a year for 24 machines. "
                "Costs are quotes and list prices; the benefit side is simulated.", h=0.6)
    rows = [["Item", "How the number is built", "Amount"],
            ["Sensor retrofit", "₹45,000 per machine × 24: wireless vibration + temperature sensor, current clamp, gateway share. Indian list prices: "
             "₹15,000 to 50,000 per wireless vibration sensor; a Banner QM30VT2 lands at ₹55,009 with duties", lakh(C["sensor_kits_inr"])],
            ["Setup and training", "Gateway and driver licence ≈ ₹1.5 L (Kepware FANUC suite $995 plus hardware); 25 engineer-days of integration at "
             "₹12,000 a day = ₹3 L; four training days and travel ≈ ₹1.5 L", lakh(C["implementation_inr"])],
            ["Cloud, per year", "Budgeted at ₹2.4 L; the estimate on the right is about ₹1.5 L, the rest is headroom", lakh(C["cloud_per_year_inr"])],
            ["Support, per year", "A ₹50,000 monthly retainer: about four engineer-days a month for model re-fits, help desk and updates", lakh(C["support_per_year_inr"])],
            ["Five-year total", f"{lakh(C['one_time_inr'])} once + {lakh(C['yearly_inr'])} × 5", lakh(D["opex_cost_inr"])]]
    table(s, LEFT, 2.0, 5.75, rows, [1.15, 3.7, 0.9], size=7.8, row_h=0.4, bold_first_col=True, align_right_cols=(2,))
    gem_in, gem_out = 40 * 0.30, 14 * 2.50
    run = 0.000018 * 2_628_000 + 0.000002 * 2 * 2_628_000
    total = gem_in + gem_out + run + 25
    rect(s, 6.25, 2.0, 3.4, 2.55, fill=LIGHT, line=None)
    text(s, 6.4, 2.08, 3.1, 0.25, "Cloud cost a month, estimated", size=9.5, bold=True)
    bullets(s, 6.4, 2.38, 3.15, 2.2, [
        f"Gemini 3.5 Flash-Lite: about 300 agent and copilot calls a day × 6k tokens ≈ 54 M tokens a month. 40 M in × $0.30 + 14 M out × $2.50 ≈ ${gem_in + gem_out:.0f} ≈ ₹{(gem_in + gem_out) * FX / 1000:.1f}k",
        f"Cloud Run, 1 vCPU and 2 GiB kept warm around the clock so the plant twin never sleeps: 2.63 M s × ($0.000018 + 2 × $0.000002) ≈ ${run:.0f} ≈ ₹{run * FX / 1000:.1f}k",
        f"Supabase Pro $25 ≈ ₹{25 * FX / 1000:.1f}k. BigQuery sandbox, Cloud Storage and Pub/Sub stay inside their free tiers",
        [(f"About ${total:.0f} ≈ ₹{total * FX / 1000:.1f}k a month, ₹{total * FX * 12 / 1e5:.1f} L a year", {"bold": True})]], size=8, gap=4)
    text(s, LEFT, 4.82, WIDTH, 0.4, [[
        ("Against ", {}), (lakh(D["opex_cost_inr"]), {"bold": True}), (" of cost the default simulation returns ", {}), (lakh(D["total_benefit_inr"]), {"bold": True}),
        (f" of benefit, {lakh(D['net_benefit_inr'])} net, over five years. Return = net ÷ cost = {D['roi_multiple']:.1f}×, payback in month {D['payback_months']}. "
         f"The pessimistic run (60% adherence, 70% coverage) nets ₹{SENS[2][2]:.1f} Cr, {SENS[2][3]:.1f}×.", {})]], size=9.5, spacing=0)
    footnote(s, f"Gemini API pricing page (Sept 2026 update); Cloud Run instance-based pricing; Supabase Pro plan; sensor prices from Zbotic (2026), tanotis.com (Banner QM30VT2) "
                f"and IndiaMART listings; day rates are our assumptions. ₹{FX} per USD. Factory economics are editable under Settings → Factory economics.", y=5.24)
    notes(s, "Support at ₹6 L a year is a full service retainer, deliberately generous. Even so the simulated payback is month 3.")


def slide_snapshots(s):
    headline(s, "Gemini at work in the prototype. It adds the procedure, the language and the judgement; the engine supplies the numbers, "
                "and every answer carries a Gemini or rule-based label.", h=0.6)
    picture(s, crop("22_incident_room_safety.png", 700, 790, 2190, 1600), LEFT, 1.85, 3.6, 2.0)
    text(s, LEFT, 3.85, 3.6, 0.45, "Incident room, third agent: the Safety officer adds what the engine cannot, lockout-tagout, a 30-minute cool-down "
         "and a burn hazard, and names the SOPs it read.", size=7.5, color=MUTED, spacing=0)
    picture(s, crop("20_briefing_gemini.png", 530, 1030, 2860, 1500), LEFT, 4.32, 3.6, 0.7)
    text(s, LEFT, 5.05, 3.6, 0.3, "Shift briefing written by Gemini from the live overview, read aloud on request.", size=7.5, color=MUTED, spacing=0)
    picture(s, crop("21_copilot_hindi.png", 560, 740, 2410, 1290), 4.2, 1.85, 5.45, 1.45)
    text(s, 4.2, 3.3, 5.45, 0.3, "Ask OpEx, in Hindi, 'what happens if I shut M-17 down now?': Gemini calls the pricing tool, quotes ₹1.98 L against "
         "₹3.99 L, and cites the lockout-tagout SOP.", size=7.5, color=MUTED, spacing=0)
    picture(s, crop("11_rewind.png", 540, 330, 2830, 1080), 4.2, 3.72, 5.45, 1.3)
    text(s, 4.2, 5.05, 5.45, 0.3, "Rewind the factory: one of the plant's own breakdowns replayed with OpEx watching, same bearing, same luck, "
         "₹8.4 L protected.", size=7.5, color=MUTED, spacing=0)
    notes(s, "Gemini screens captured 18 Oct 2026 with an AI Studio key, plus the Rewind page. Also in the app: investigation reports and the part-photo check.")


def slide_performance(s):
    headline(s, "Model scores on held-out data, and a simulated plant result that stays inside published ranges even when we make "
                "our own assumptions worse", h=0.6)
    rows = [["Model (held-out test)", "Data", "Result"],
            ["Component precursor model, 48 h horizon", "Azure PdM telemetry (synthetic sample), trained Jan-Aug 2015, tested Sep-Dec",
             "Caught 94 to 100% of component failures 24 to 32 h ahead; 11.9 false alarms per machine-year (1.2 to 5.8 per component), all charged in the simulator"],
            ["Remaining-life model + 10/90% quantiles", "NASA C-MAPSS FD001, NASA's 100 test engines",
             f"RMSE {RUL['test_rmse_cycles']:.1f} cycles (published FD001 results sit around 11 to 16); 10-90% band covers {RUL['p10_p90_coverage']:.0%}; every engine warned ≥ 10 cycles ahead"],
            ["Failure-mode classifier", "AI4I 2020 (synthetic, known to be easy), 2,000 test rows",
             f"ROC AUC {FM['roc_auc_any_failure']:.3f}, precision {FM['precision_at_threshold']:.0%}, recall {FM['recall_at_threshold']:.0%}; tool-wear failures are near-random in this set (11% caught)"],
            ["App", "24 machines, 1-minute data", "Overview 0.18 s warm; 5-year simulation (2.1 M machine-hours) about 4.5 s; 38 automated tests pass"]]
    table(s, LEFT, 1.95, 5.6, rows, [1.5, 1.5, 2.6], size=7.4, row_h=0.46, bold_first_col=True)
    data = CategoryChartData()
    data.categories = ["Fewer breakdowns", "Lower maintenance cost"]
    data.add_series("OpEx, simulated", (round(D["breakdowns_reduction"] * 100, 1), round(D["maintenance_cost_reduction"] * 100, 1)))
    data.add_series("US DOE guide, low", (70, 25))
    data.add_series("US DOE guide, high", (75, 30))
    data.add_series("Deloitte", (70, 25))
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.0), Inches(2.1), Inches(3.65), Inches(2.5), data)
    style_chart(frame.chart, [BLUE, "BDC1C6", "80868B", "DADCE0"])
    text(s, 6.0, 1.9, 3.65, 0.2, "Simulated improvement vs published ranges (a check, not a result)", size=8, bold=True)
    bullets(s, LEFT, 4.66, 9.25, 0.6, [
        [("Validation status. ", {"bold": True}), ("Not yet run on a real plant. The live plant is a calibrated twin and the breakdown reduction follows from our own "
          "prevention assumption (recall × 90% coverage × 85% adherence), so the chart checks that assumption against published ranges; it does not validate it. "
          "The next slide shows the result at 70%, 60% and 50% adherence.", {})],
        [("Downtime fell 63% against the DOE's 35-45% ", {"bold": True}), ("because ours counts equipment stops only; material and set-up stops (17.5% of planned "
          "time, from a real line) are identical in both worlds, which is also why OEE moves just +1.0 point.", {})]], size=7.8, gap=2)
    footnote(s, "Ranges: US DOE/FEMP O&M Best Practices Guide R3.0 §5.4 ('independent surveys'); Deloitte Analytics Institute position paper. Scores: models/*.json.", y=5.3)
    notes(s, "Lead with Azure PdM and C-MAPSS; AI4I is the weak dataset and we say so. The chart is a calibration check, not a result.")


def slide_calc(s, donor):
    give_title(s, donor, "How we calculated the 5-year numbers")
    headline(s, f"Same plant, same luck, two policies. Every random draw is shared, so the {lakh(D['net_benefit_inr'])} comes only from the decisions.")
    rows = [["Benefit line", "Formula and inputs", "5 years"],
            ["Production protected", f"{D['good_parts_gained']:,.0f} extra good parts × contribution ₹175 to 760 per part", lakh(D["production_protected_inr"])],
            ["Maintenance saved", "parts + technician hours at ₹300/h (× 1.5 after a breakdown) + secondary damage (40% chance of a ₹2.2 L spindle job) + inspections", lakh(D["maintenance_savings_inr"])],
            ["Scrap avoided", f"{D['defects_avoided']:,.0f} fewer bad parts × ₹380 to 1,450 scrap cost", lakh(D["scrap_savings_inr"])],
            ["Energy saved", f"{D['energy_saved_mwh']:,.0f} MWh × ₹9.25/kWh (MSEDCL HT-I(A), MERC order 25 Mar 2026); {D['co2_avoided_t']:,.0f} t CO₂ at 0.675 t/MWh", lakh(D["energy_savings_inr"])],
            ["Total benefit", "sum of the four lines", lakh(D["total_benefit_inr"])],
            ["Net and return", f"{lakh(D['total_benefit_inr'])} − {lakh(D['opex_cost_inr'])} cost = {lakh(D['net_benefit_inr'])}; ÷ cost = {D['roi_multiple']:.1f}×; payback in month {D['payback_months']}", f"{D['roi_multiple']:.1f}×"]]
    table(s, LEFT, 1.85, 5.7, rows, [1.25, 3.6, 0.85], size=7.6, row_h=0.3, bold_first_col=True, align_right_cols=(2,))
    text(s, LEFT, 4.08, 5.7, 0.22, "Breakdowns a year, without and with OpEx", size=8, bold=True)
    data = CategoryChartData()
    data.categories = [f"Year {y['year']}" for y in N["yearly"]]
    data.add_series("Without OpEx", [y["without"]["failures"] for y in N["yearly"]])
    data.add_series("With OpEx", [y["with"]["failures"] for y in N["yearly"]])
    frame = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(LEFT), Inches(4.28), Inches(5.7), Inches(1.05), data)
    style_chart(frame.chart, ["9AA0A6", BLUE], label_fmt="0", size=7)
    frame.chart.legend.position = XL_LEGEND_POSITION.RIGHT
    text(s, 6.3, 1.85, 3.35, 0.22, "Where the inputs come from", size=9, bold=True)
    bullets(s, 6.3, 2.1, 3.35, 1.55, [
        "Failures per machine-year: bearing 1.79, drive 2.59, electrical 1.92, coolant 1.31 (Azure PdM, 761 failures)",
        "Life spread Weibull k = 4.41; wear curve damage(u) = (e^3.75u − 1)/(e^3.75 − 1) (NASA C-MAPSS)",
        "Prevented = model recall × sensor coverage 0.90 × adherence 0.85 ≈ 72 to 76%; warning error N(0.04, 0.20)",
        "Minor stops 8.9%, set-ups 8.6%, repair-time σ 0.82, performance 92.9%, quality 98.3% (Zenodo 17855209, a real line at 59% OEE)",
        "6,750 planned hours a year (3 shifts × 7.5 h × 300 days)"], size=7.5, gap=2)
    text(s, 6.3, 3.66, 3.35, 0.22, "If people trust it less and sensors see less", size=9, bold=True)
    srows = [["Adherence", "Coverage", "Net 5 yr", "Return", "Payback"]]
    for a, c, net, roi, pb, br in SENS[:4]:
        srows.append([f"{a:.0%}", f"{c:.0%}", f"₹{net:.1f} Cr", f"{roi:.1f}×", f"month {pb}"])
    table(s, 6.3, 3.9, 3.35, srows, [0.72, 0.7, 0.72, 0.56, 0.65], size=7.5, row_h=0.23, align_right_cols=(2, 3, 4))
    text(s, 6.3, 5.14, 3.35, 0.3, "Same seed and plant; only the two assumptions change. The baseline of 7.6 failures per machine-year is Azure PdM's; "
         "calmer fleets sit lower in this table.", size=7, color=MUTED, spacing=0)
    notes(s, "Walk the table, then the sensitivity block: the case does not depend on everyone trusting the AI.")


def slide_future(s):
    headline(s, "What comes next, why it can grow, and who gains besides the plant's accountant")
    cols = [
        ("Next steps", [
            [("Pilot. ", {"bold": True}), ("Six machines in one plant with an OPC UA / MTConnect gateway; health estimators re-fitted per machine during commissioning. "
                                            "No plant has signed yet; the ask of the finale is introductions to two ACMA members.", {})],
            [("6 to 12 months. ", {"bold": True}), ("Multi-plant view and benchmarking; Dataflow into BigQuery; per-machine fine-tuning on Vertex AI once six months of real data exist; "
                                                     "Live API voice for operators.", {})],
            [("Later. ", {"bold": True}), ("Line-camera inspection with Visual Inspection AI; spare-parts auto-ordering; tariff-aware scheduling.", {})]]),
        ("Why it can grow", [
            [("The 1,100+ ACMA members alone, at ₹8.4 lakh a year each, are about ₹92 crore a year of recurring revenue, before counting non-member job shops. "
              "India bought about USD 3.7 billion of machine tools in FY25, roughly half for automotive (IMTMA).", {})],
            [("The models see health-index features rather than raw machine-specific sensors, so they transfer to any machine that can produce a health index. "
              "One container per plant today; the same image serves a plant group once readings are real, with Postgres holding the state.", {})],
            [("Cloud costs about ₹1.5 lakh a year per plant at full use, and nothing until a plant decides to pay.", {})]]),
        ("Who else gains", [
            [("Workers. ", {"bold": True}), ("1.5 million people work directly in Indian auto components (ACMA). Fewer 2 a.m. emergency repairs; lockout-tagout steps are part of "
                                              "every stop recommendation.", {})],
            [("Language and access. ", {"bold": True}), ("Hindi, Telugu, Tamil, Kannada, Marathi; read-aloud and voice for low-literacy and low-vision staff; WCAG 2.2 AA.", {})],
            [("Energy. ", {"bold": True}), (f"{D['energy_saved_mwh']:,.0f} MWh and {D['co2_avoided_t']:,.0f} t CO₂ avoided per 24-machine plant over five years in the simulation, "
                                             "mostly from idle spindles and coolant pumps put on standby.", {})]]),
    ]
    for i, (title, items) in enumerate(cols):
        x = LEFT + i * 3.15
        text(s, x, 1.95, 2.95, 0.25, title, size=10.5, bold=True)
        hline(s, x, 2.25, 2.95)
        text(s, x, 2.35, 2.95, 2.9, items, size=9.5, spacing=7, line=1.1)
    footnote(s, "IMTMA figure as reported at the AMT/IMTMA briefing, 2025; ACMA FY26 review. Revenue figure is members × our list price, not a forecast.", y=5.25)
    notes(s, "Be plain about the pilot status. The bottom-up revenue number is more credible to a judge than an analyst market size.")


def slide_links(s):
    box = title_box(s)
    links = {"GitHub Public Repository": "github.com/srihitha2005/Opex",
             "Demo Video Link (3 Minutes)": "[add before submitting]",
             "Final Product Link": "[add before submitting]"}
    for p in box.text_frame.paragraphs:
        label = "".join(r.text for r in p.runs).strip()
        if label in links and p.runs:
            r = copy.deepcopy(p.runs[0]._r)
            p._p.append(r)
            p.runs[-1].text = ":  " + links[label]
            p.runs[-1].font.color.rgb = rgb(BLUE)
    text(s, LEFT, 2.45, 5.5, 0.3, "Try it in 60 seconds", size=11, bold=True)
    steps = ["Click 'Explore as a guest' on the sign-in page (demo accounts are listed in the README).",
             "Open M-17. Read steps 1 to 4, open the AI incident room, compare the priced options, approve one.",
             "Open 5-year simulation, choose a 20-second replay, Simulate, Skip to results, then Rewind the factory.",
             "Ask OpEx: 'What happens if I shut down Machine 17?' by voice, in Hindi if you like.",
             "Settings → Display: high contrast, extra-large text, Lexend, simple words."]
    text(s, LEFT, 2.8, 5.5, 2.3, [[(f"{i + 1}.  ", {"bold": True}), (t, {})] for i, t in enumerate(steps)], size=9.5, spacing=6, line=1.06)
    picture(s, crop("09_simulation_live.png", 540, 300, 1680, 1170), 6.15, 2.45, 3.5, 2.65)
    text(s, 6.15, 5.12, 3.5, 0.3, "5-year simulation at day 665: value created so far and the running with-vs-without table.", size=7.5, color=MUTED)
    text(s, LEFT, 5.12, 5.5, 0.25, "AUDIT.md explains every key in plain language; HOW_TO_RUN.md has the 3-minute tour.", size=7.5, color=MUTED)
    notes(s, "Replace the two bracketed items with the real links before exporting.")


def slide_sources(s, donor):
    give_title(s, donor, "Sources: every external number on these slides, and where it comes from")
    left = [
        [("1  ACMA, ", {"bold": True}), ("Industry Performance Review FY2025-26 and acma.in/about-us: turnover ₹7.59 lakh crore (USD 85.9 bn), 1,100+ members, 1.5 million direct jobs, 2.3% of GDP.", {})],
        [("2  ABB, ", {"bold": True}), ("'Value of Reliability' survey (Sapio Research, July 2023, n = 3,215, India cut): 88% see unplanned outages at least monthly; 19% run to fail; about ₹7 million lost per hour. new.abb.com/news/detail/108271", {})],
        [("3  US DOE / FEMP, ", {"bold": True}), ("O&M Best Practices Guide Release 3.0 (PNNL-19634, 2010), §5.4: predictive maintenance 'industrial average savings' of breakdowns −70-75%, maintenance cost −25-30%, downtime −35-45%, ROI 10×. The guide attributes these to independent surveys; we use them only as a sanity range.", {})],
        [("4  Deloitte Analytics Institute, ", {"bold": True}), ("Predictive Maintenance position paper: breakdowns −70%, maintenance costs −25%, productivity +25%, uptime +10-20%.", {})],
        [("5  Datasets. ", {"bold": True}), ("AI4I 2020 (UCI 601, CC BY 4.0; synthetic); NASA C-MAPSS FD001; Microsoft Azure Predictive Maintenance sample (synthetic); Zenodo record 17855209 (real serial line, 62 shifts, 635 stops, OEE 59.1% before improvement, CC BY 4.0); NIST Nestor vocabulary.", {})],
        [("6  Electricity. ", {"bold": True}), ("MERC order of 25 March 2026 (Case 75 of 2025) on MSEDCL tariffs; HT-I(A) energy plus wheeling entered as ₹9.25/kWh in backend/app/plant/factory.py and editable in Settings.", {})],
        [("7  Grid emissions. ", {"bold": True}), ("CEA CO₂ Baseline Database v22 (Aug 2026) weighted average 0.675 t/MWh as entered in the app; the v21 guide published 0.710 for FY2024-25 and v20 0.727 for FY2023-24. cea.nic.in", {})],
        [("8  Wages and repairs. ", {"bold": True}), ("In-house technician about ₹19k a month (Indeed India, Oct 2026) blended with OEM visits to ₹300/h; engineer day rate ₹12,000 is our assumption; BHEL spindle-bearing repair tender (2020) ₹2.88 L sets the scale of secondary damage.", {})],
        [("9  Plant economics. ", {"bold": True}), ("Contribution ₹175 to 760 per part (₹4,200 to 6,500 per machine-hour), scrap ₹380 to 1,450, repair hours and stock-out odds are the demo plant's assumptions in factory.py, all editable under Settings → Factory economics.", {})],
        [("10  Model literature. ", {"bold": True}), ("C-MAPSS FD001 RMSE of 11 to 16 cycles is the range of commonly cited CNN, LSTM and transformer results; our 13.0 is in that range, not above it.", {})],
        [("11  Standards. ", {"bold": True}), ("ISA-101 HMI design, ISA-18.2 alarm management, ISO 10816-3 vibration zones (heuristic for spindles), WCAG 2.2 AA.", {})],
    ]
    right = [
        [("12  Gemini API pricing, ", {"bold": True}), ("ai.google.dev/gemini-api/docs/pricing (updated Sept 2026): gemini-3.5-flash-lite $0.30 in / $2.50 out per 1M tokens; gemini-3.8-flash $0.75 / $3.75 (to 31 Dec 2026); gemini-3.8-live audio $3.00 / $12.00. The embedding model we call, gemini-embedding-001, is listed as superseded by gemini-embedding-2 at $0.20 per 1M; switching is a one-line model-name change. Free tier available for the Flash models; daily limits are per project and shown in AI Studio.", {})],
        [("13  Cloud Run pricing, ", {"bold": True}), ("cloud.google.com/run/pricing: instance-based $0.000018 per vCPU-second and $0.000002 per GiB-second; free tier 180,000 vCPU-seconds, 360,000 GiB-seconds and 2 million requests a month.", {})],
        [("14  Other Google Cloud prices. ", {"bold": True}), ("Vertex AI Agent Engine runtime $0.0864 per vCPU-hour, $0.009 per GiB-hour (2026); Visual Inspection AI $100 per camera stream per solution per month (cloud.google.com/vision-ai/pricing); Speech-to-Text v2 $0.016 per minute; Dataflow streaming about $0.069 per vCPU-hour, $0.003557 per GiB-hour (cloud.google.com/dataflow/pricing); Vector Search about $0.75 per node-hour for e2-standard-16 (Google developer forum estimate); Looker Studio Pro $9 per user per project per month.", {})],
        [("15  Supabase, ", {"bold": True}), ("Pro plan $25 a month; free projects pause after 7 days without use. supabase.com/pricing", {})],
        [("16  Sensor and gateway prices. ", {"bold": True}), ("Zbotic (2026 guide): commercial wireless vibration sensors ₹15,000 to 50,000; Banner QM30VT2 ₹55,009 incl. duties (tanotis.com); IndiaMART single-axis wireless vibration sensor ₹12,500; Kepware FANUC FOCAS driver suite about $995 (industrialmonitordirect.com).", {})],
        [("17  Market. ", {"bold": True}), ("IMTMA: Indian machine-tool consumption about USD 3.7 bn in FY2024-25, about half from automotive (AMT/IMTMA briefing, 2025). The ₹92 crore figure is 1,100 members × ₹8.4 lakh a year.", {})],
        [("18  Exchange rate. ", {"bold": True}), (f"₹{FX} per USD, interbank close of 17 Oct 2026 (reported between 94.71 and 94.96).", {})],
        [("19  Simulation. ", {"bold": True}), ("Default run of the app's 5-year simulation, seed 2026 (pitch/numbers.json); sensitivity runs in pitch/sensitivity.json; model scores in models/*.json; every simulator input with its source and method in data/processed/calibration.json.", {})],
    ]
    text(s, LEFT, 1.25, 4.55, 4.2, left, size=7, spacing=3.5, line=1.05)
    text(s, 5.1, 1.25, 4.55, 4.2, right, size=7, spacing=3.5, line=1.05)
    notes(s, "Everything with a number has a line here. Where a figure could not be checked against a primary document, the line says so.")


def build() -> Path:
    prs = Presentation(PITCH / "template.pptx")
    sl = list(prs.slides)
    slide_team(sl[0])
    slide_brief(sl[2])
    slide_opportunities(sl[3])
    slide_features(sl[4])
    slide_flow(sl[5])
    slide_wireframes(sl[6])
    slide_architecture(sl[7])
    slide_tech(sl[8])
    slide_cost(sl[9])
    slide_snapshots(sl[10])
    slide_performance(sl[11])
    slide_future(sl[12])
    slide_links(sl[13])
    slide_calc(sl[14], sl[2])
    slide_tech_optional(sl[16], sl[2])
    slide_sources(sl[17], sl[2])
    notes(sl[15], "Thank you. Team OpEx.")

    # Drop the template's instruction slide (#2); order the rest.
    id_list = prs.slides._sldIdLst
    entries = list(id_list)
    instruction = entries[1]
    prs.part.drop_rel(instruction.rId)
    id_list.remove(instruction)
    order = [0, 2, 3, 4, 5, 6, 7, 8, 16, 9, 10, 11, 14, 12, 13, 17, 15]
    by_original = {idx: e for idx, e in enumerate(entries)}
    for e in list(id_list):
        id_list.remove(e)
    for idx in order:
        id_list.append(by_original[idx])

    out = PITCH / "OpEx_AI_Builder_Cup_2026.pptx"
    prs.save(out)
    return out


if __name__ == "__main__":
    print(build())
